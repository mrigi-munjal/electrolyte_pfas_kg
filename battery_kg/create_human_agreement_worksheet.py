#!/usr/bin/env python3
"""
Create a blinded human-annotation worksheet for the LLM-as-Judge evaluation.

Samples 50 triples from the 600 already judged by Mistral-Large-3-675B and
LLaMA-3.1-405B (evaluation_combined_*.csv) so that a human annotator can be
added as a third rater.

Sampling is stratified on the joint judge outcome (PASS/PASS, FAIL/FAIL,
FAIL/PASS, PASS/FAIL) and balanced across the two datasets. The rare cells
(where the judges disagree, or agree on FAIL) are deliberately oversampled --
a plain random 50 would contain ~3 judge-disagreement cases and be useless for
saying anything about them. Inclusion probabilities are recorded in the answer
key so the analysis script can weight the sample back to the full 600.

The worksheet itself is blinded: no judge labels, shuffled row order.

Input:  evaluation_output/evaluation_combined_{timestamp}.csv
Output: evaluation_output/human_worksheet_{timestamp}.xlsx   (annotate this)
        evaluation_output/human_worksheet_{timestamp}.csv     (same, CSV)
        evaluation_output/human_worksheet_key_{timestamp}.csv (do NOT open
                                                               before annotating)
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "evaluation_output"

# Joint judge cell -> number of triples to sample. Sums to 50.
CELL_QUOTAS = {
    ("PASS", "PASS"): 20,
    ("FAIL", "FAIL"): 10,
    ("FAIL", "PASS"): 10,
    ("PASS", "FAIL"): 10,
}

RANDOM_SEED = 42

# Same rubric the LLM judges were given, restated for a human annotator.
INSTRUCTIONS = [
    "HUMAN ANNOTATION WORKSHEET - knowledge graph triple correctness",
    "",
    "For each row, read the SOURCE TEXT and the extracted triple",
    "(SUBJECT, RELATION, OBJECT), then fill in the human_correctness column",
    "with either PASS or FAIL.",
    "",
    "RUBRIC (identical to the one given to the LLM judges):",
    "",
    "  Is the triple consistent with the source text? It does not need to be a",
    "  verbatim extraction -- paraphrasing, reasonable inference, and standard",
    "  domain knowledge are acceptable. Mark FAIL only if the triple clearly",
    "  contradicts or cannot be supported by the source text.",
    "",
    "Use human_notes for anything ambiguous. Leave human_correctness blank only",
    "if the row is genuinely unjudgeable; blanks are dropped from the analysis.",
    "",
    "The rows are shuffled and the LLM judgements are hidden on purpose -- do",
    "not look at human_worksheet_key_*.csv until you have finished annotating.",
]


def find_latest_evaluation_file() -> Path:
    files = list(OUTPUT_DIR.glob("evaluation_combined_*.csv"))
    if not files:
        print("ERROR: no evaluation_combined_*.csv found in evaluation_output/")
        sys.exit(1)
    return max(files, key=lambda f: f.stat().st_mtime)


def stratified_sample(df: pd.DataFrame, seed: int) -> pd.DataFrame:
    """Sample each judge cell, splitting the quota evenly across datasets."""
    rng = np.random.RandomState(seed)
    picked = []

    for (mistral, llama), quota in CELL_QUOTAS.items():
        cell = df[
            (df["mistral_correctness"] == mistral)
            & (df["llama_correctness"] == llama)
        ]
        if len(cell) == 0:
            print(f"WARNING: judge cell {mistral}/{llama} is empty, skipping")
            continue

        datasets = sorted(cell["dataset"].unique())
        per_dataset = {d: quota // len(datasets) for d in datasets}
        for d in datasets[: quota % len(datasets)]:
            per_dataset[d] += 1

        # A dataset that cannot fill its share hands the remainder back.
        deficit = 0
        chosen = []
        for d in datasets:
            sub = cell[cell["dataset"] == d]
            take = min(per_dataset[d], len(sub))
            deficit += per_dataset[d] - take
            chosen.append(sub.sample(n=take, random_state=rng))
        if deficit:
            already = pd.concat(chosen).index
            spare = cell.drop(index=already)
            take = min(deficit, len(spare))
            if take:
                chosen.append(spare.sample(n=take, random_state=rng))
            if take < deficit:
                print(
                    f"WARNING: judge cell {mistral}/{llama} only has {len(cell)} "
                    f"triples, sampled all of them instead of {quota}"
                )

        cell_sample = pd.concat(chosen)
        # Inclusion probability within this cell -> design weight for analysis.
        cell_sample = cell_sample.assign(
            judge_cell=f"{mistral}/{llama}",
            cell_population=len(cell),
            cell_sampled=len(cell_sample),
            design_weight=len(cell) / len(cell_sample),
        )
        picked.append(cell_sample)

    return pd.concat(picked)


def write_xlsx(path: Path, worksheet: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame({"instructions": INSTRUCTIONS}).to_excel(
            writer, sheet_name="Instructions", index=False
        )
        worksheet.to_excel(writer, sheet_name="Annotation", index=False)

        instr = writer.sheets["Instructions"]
        instr.column_dimensions["A"].width = 80
        instr["A1"].font = Font(bold=True)

        sheet = writer.sheets["Annotation"]
        widths = {
            "triple_id": 10,
            "subject": 28,
            "relation": 22,
            "object": 34,
            "source_text": 90,
            "human_correctness": 18,
            "human_notes": 34,
        }
        for idx, col in enumerate(worksheet.columns, start=1):
            letter = get_column_letter(idx)
            sheet.column_dimensions[letter].width = widths.get(col, 18)
            sheet[f"{letter}1"].font = Font(bold=True)

        text_col = get_column_letter(list(worksheet.columns).index("source_text") + 1)
        for row in range(2, len(worksheet) + 2):
            sheet.row_dimensions[row].height = 90
            sheet[f"{text_col}{row}"].alignment = Alignment(
                wrap_text=True, vertical="top"
            )

        verdict_col = get_column_letter(
            list(worksheet.columns).index("human_correctness") + 1
        )
        dv = DataValidation(
            type="list", formula1='"PASS,FAIL"', allow_blank=True, showDropDown=False
        )
        sheet.add_data_validation(dv)
        dv.add(f"{verdict_col}2:{verdict_col}{len(worksheet) + 1}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create a blinded 50-triple human annotation worksheet"
    )
    parser.add_argument("--input", "-i", default=None, help="evaluation_combined CSV")
    parser.add_argument("--seed", type=int, default=RANDOM_SEED)
    args = parser.parse_args()

    input_file = Path(args.input) if args.input else find_latest_evaluation_file()
    if not input_file.exists():
        print(f"ERROR: input file not found: {input_file}")
        sys.exit(1)

    df = pd.read_csv(input_file)
    print(f"Input file: {input_file}")
    print(f"Loaded {len(df)} judged triples")

    # The judges never returned ERROR on this run, but guard anyway.
    df = df[
        df["mistral_correctness"].isin(["PASS", "FAIL"])
        & df["llama_correctness"].isin(["PASS", "FAIL"])
    ].reset_index(drop=True)

    sample = stratified_sample(df, args.seed)
    sample = sample.sample(frac=1.0, random_state=args.seed).reset_index(drop=True)
    sample["triple_id"] = [f"T{i:03d}" for i in range(1, len(sample) + 1)]

    print(f"\nSampled {len(sample)} triples")
    print(
        sample.groupby(["judge_cell", "dataset"])
        .size()
        .unstack(fill_value=0)
        .to_string()
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    worksheet = sample[
        ["triple_id", "subject", "relation", "object", "source_text"]
    ].copy()
    worksheet["human_correctness"] = ""
    worksheet["human_notes"] = ""

    xlsx_path = OUTPUT_DIR / f"human_worksheet_{timestamp}.xlsx"
    csv_path = OUTPUT_DIR / f"human_worksheet_{timestamp}.csv"
    key_path = OUTPUT_DIR / f"human_worksheet_key_{timestamp}.csv"

    write_xlsx(xlsx_path, worksheet)
    worksheet.to_csv(csv_path, index=False)

    sample[
        [
            "triple_id",
            "dataset",
            "subject",
            "relation",
            "object",
            "mistral_correctness",
            "llama_correctness",
            "agreement",
            "judge_cell",
            "cell_population",
            "cell_sampled",
            "design_weight",
        ]
    ].to_csv(key_path, index=False)

    print(f"\nWorksheet (annotate this): {xlsx_path}")
    print(f"Worksheet CSV:             {csv_path}")
    print(f"Answer key (do not peek):  {key_path}")
    print(
        "\nWhen the human_correctness column is filled in, run:\n"
        f"  python battery_kg/analyze_human_agreement.py --worksheet {xlsx_path}"
    )


if __name__ == "__main__":
    main()
